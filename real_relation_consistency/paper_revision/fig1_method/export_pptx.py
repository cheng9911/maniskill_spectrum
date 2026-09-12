from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Mm,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
OUT=Path(__file__).resolve().parent
d=json.loads((OUT/'drawing_operations.json').read_text());prs=Presentation();prs.slide_width=Mm(d['width_mm']);prs.slide_height=Mm(d['height_mm']);sl=prs.slides.add_slide(prs.slide_layouts[6])
def rgb(s):return RGBColor.from_string(s.lstrip('#'))
for o in d['operations']:
 k=o['kind']
 if k=='text':
  w=o['w'];x=o['x']-({'left':0,'center':w/2,'right':w}[o['align']]);h=o['h'];y=o['y']-h/2
  sh=sl.shapes.add_textbox(Mm(x),Mm(y),Mm(w),Mm(h));tf=sh.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
  for i,t in enumerate(o['text'].split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[o['align']];p.space_before=p.space_after=Pt(0);r=p.add_run();r.text=t;r.font.name='DejaVu Sans';r.font.size=Pt(o['size']);r.font.bold=o['bold'];r.font.color.rgb=rgb(o['color'])
 elif k in ['rect','ellipse']:
  sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE if k=='rect' else MSO_SHAPE.OVAL,Mm(o['x']),Mm(o['y']),Mm(o['w']),Mm(o['h']));sh.fill.solid();sh.fill.fore_color.rgb=rgb(o['fill']);sh.line.color.rgb=rgb(o['edge']);sh.line.width=Pt(o['lw'])
 else:
  pts=o['points'];fb=sl.shapes.build_freeform(int(Mm(pts[0][0])),int(Mm(pts[0][1])));fb.add_line_segments([(int(Mm(x)),int(Mm(y))) for x,y in pts[1:]],close=False);sh=fb.convert_to_shape();sh.fill.background();sh.line.color.rgb=rgb(o['color']);sh.line.width=Pt(o['lw'])
  if o['dash']:sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
  if o['arrow']:
   tail=OxmlElement('a:tailEnd');tail.set('type','triangle');tail.set('w','sm');tail.set('len','sm');sh.line._get_or_add_ln().append(tail)
sl.notes_slide.notes_text_frame.text='All shapes, labels and curve paths are editable. Source-response curves use archived Pdiag finite data; other paths/geometry are schematic. See provenance.json and figure caption. No hardware results illustrated.'
prs.save(OUT/'fig1_method_editable.pptx')
assert len(sl.shapes)==len(d['operations'])
print('Saved',len(sl.shapes),'editable shapes; no slide screenshots')
